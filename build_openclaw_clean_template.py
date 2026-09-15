from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


WORKDIR = Path(r"D:\Desktop\Projects\PPT模板替换\openclaw_rebuild")
SOURCE_PPT = WORKDIR / "source.pptx"
SLIDES_DIR = WORKDIR / "slides"
ASSETS_DIR = Path(r"D:\Desktop\Projects\PPT模板替换\template_assets")
OUTPUT_PPT = WORKDIR / "openclaw_clean_templated.pptx"

BG_COVER = ASSETS_DIR / "image5.jpeg"
BG_CONTENT = ASSETS_DIR / "image5.jpeg"


def sorted_slide_images() -> list[Path]:
    def sort_key(path: Path) -> int:
        match = re.search(r"(\d+)", path.stem)
        return int(match.group(1)) if match else 9999

    return sorted(SLIDES_DIR.glob("*.PNG"), key=sort_key)


def read_cover_text() -> tuple[str, str]:
    prs = Presentation(str(SOURCE_PPT))
    slide = prs.slides[0]
    texts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                texts.append(text)
    title = texts[0] if texts else "OpenClaw 探索"
    subtitle = texts[1] if len(texts) > 1 else ""
    return title, subtitle


def add_full_bg(slide, image_path: Path, slide_width, slide_height) -> None:
    slide.shapes.add_picture(str(image_path), 0, 0, width=slide_width, height=slide_height)


def add_textbox(slide, left, top, width, height, text, *, font_size, bold=False, color=(0, 0, 0), font_name="Microsoft YaHei", align=None):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return tx


def build_cover(slide, title: str, subtitle: str, cover_shot: Path, slide_width, slide_height) -> None:
    add_full_bg(slide, BG_COVER, slide_width, slide_height)

    # Soft white panel so the long title stays readable.
    panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(1.05),
        Inches(7.15),
        Inches(4.9),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    panel.fill.transparency = 0.12
    panel.line.color.rgb = RGBColor(255, 255, 255)
    panel.line.transparency = 1

    kicker = slide.shapes.add_textbox(Inches(0.88), Inches(1.35), Inches(2.1), Inches(0.35))
    p = kicker.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "OpenClaw Special"
    r.font.name = "Arial"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(246, 106, 42)

    title_box = slide.shapes.add_textbox(Inches(0.85), Inches(1.72), Inches(6.2), Inches(2.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title.replace("：", "：\n", 1) if "：" in title else title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = RGBColor(17, 24, 39)

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.88), Inches(4.85), Inches(2.5), Inches(0.45))
        p = sub.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(18)
        r.font.bold = False
        r.font.color.rgb = RGBColor(59, 130, 246)

    shot_frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(7.6),
        Inches(1.35),
        Inches(4.55),
        Inches(3.0),
    )
    shot_frame.fill.solid()
    shot_frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shot_frame.line.color.rgb = RGBColor(210, 226, 242)
    shot_frame.line.width = Pt(1.25)
    slide.shapes.add_picture(str(cover_shot), Inches(7.78), Inches(1.53), width=Inches(4.2), height=Inches(2.36))


def build_content(slide, image_path: Path, index: int, total: int, slide_width, slide_height) -> None:
    add_full_bg(slide, BG_CONTENT, slide_width, slide_height)

    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(0.62),
        Inches(10.65),
        Inches(5.98),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(219, 234, 254)
    card.line.width = Pt(1.0)

    slide.shapes.add_picture(str(image_path), Inches(0.73), Inches(0.8), width=Inches(10.3), height=Inches(5.79))

    num = slide.shapes.add_textbox(Inches(0.62), Inches(6.8), Inches(0.9), Inches(0.28))
    p = num.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"{index:02d}/{total:02d}"
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235)


def main() -> None:
    images = sorted_slide_images()
    if not images:
        raise SystemExit("No slide images found.")

    title, subtitle = read_cover_text()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    for idx, image_path in enumerate(images, start=1):
        slide = prs.slides.add_slide(blank)
        if idx == 1:
            build_cover(slide, title, subtitle, image_path, prs.slide_width, prs.slide_height)
        else:
            build_content(slide, image_path, idx, len(images), prs.slide_width, prs.slide_height)

    prs.save(str(OUTPUT_PPT))
    print(OUTPUT_PPT)


if __name__ == "__main__":
    main()
