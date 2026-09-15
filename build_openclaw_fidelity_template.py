from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


WORKDIR = Path(r"D:\Desktop\Projects\PPT模板替换\openclaw_rebuild")
SLIDES_DIR = WORKDIR / "slides"
OUTPUT_PPT = WORKDIR / "openclaw_fidelity_templated.pptx"
BG_RGB = RGBColor(249, 251, 253)


def sorted_slide_images() -> list[Path]:
    def sort_key(path: Path) -> int:
        match = re.search(r"(\d+)", path.stem)
        return int(match.group(1)) if match else 9999

    return sorted(SLIDES_DIR.glob("*.PNG"), key=sort_key)


def add_bg(slide, prs: Presentation) -> None:
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_RGB
    bg.line.fill.background()


def add_soft_frame(slide, x: float, y: float, w: float, h: float) -> None:
    shadow = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x + 0.05),
        Inches(y + 0.05),
        Inches(w),
        Inches(h),
    )
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(15, 23, 42)
    shadow.fill.transparency = 0.93
    shadow.line.fill.background()

    frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x - 0.03),
        Inches(y - 0.03),
        Inches(w + 0.06),
        Inches(h + 0.06),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor(214, 223, 233)
    frame.line.width = Pt(0.8)


def add_full_image(slide, image_path: Path, x: float, y: float, w: float, h: float) -> None:
    add_soft_frame(slide, x, y, w, h)
    slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_cover(prs: Presentation, slide, cover_image: Path) -> None:
    add_bg(slide, prs)
    add_full_image(slide, cover_image, 0.46, 0.42, 12.08, 6.79)


def add_content(prs: Presentation, slide, image_path: Path) -> None:
    add_bg(slide, prs)
    add_full_image(slide, image_path, 0.34, 0.42, 12.38, 6.96)


def main() -> None:
    images = sorted_slide_images()
    if not images:
        raise SystemExit("No slide images found.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, image in enumerate(images, start=1):
        slide = prs.slides.add_slide(blank)
        if idx == 1:
            add_cover(prs, slide, image)
        else:
            add_content(prs, slide, image)

    prs.save(str(OUTPUT_PPT))
    print(OUTPUT_PPT)


if __name__ == "__main__":
    main()
