from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any
import re

from pptx import Presentation


PT = 12700
FONT_ALIAS = {
    "微软雅黑": "Microsoft YaHei",
}


def safe_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def max_font_size_pt(shape) -> float:
    sizes = []
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                sizes.append(run.font.size / PT)
    return max(sizes) if sizes else 0.0


def collect_fonts(prs: Presentation) -> list[str]:
    fonts = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.name:
                        fonts.add(run.font.name)
    return sorted(fonts)


def get_installed_fonts() -> set[str]:
    try:
        import winreg
    except ImportError:
        return set()

    installed = set()
    keys = [
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
        (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
    ]
    for root, path in keys:
        try:
            with winreg.OpenKey(root, path) as key:
                count = winreg.QueryInfoKey(key)[1]
                for idx in range(count):
                    name, _, _ = winreg.EnumValue(key, idx)
                    cleaned = name.split("(")[0].strip()
                    if cleaned:
                        installed.add(cleaned)
        except OSError:
            continue
    return installed


def normalize_font_name(font_name: str) -> str:
    normalized = FONT_ALIAS.get(font_name.strip(), font_name.strip())
    normalized = normalized.lower()
    normalized = normalized.replace("_gbk", "")
    normalized = re.sub(r"\b(bold|light|regular|italic|ui)\b", "", normalized)
    normalized = normalized.replace("&", " ")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def font_is_available(font_name: str, installed_fonts: set[str]) -> bool:
    if not installed_fonts:
        return True
    target = normalize_font_name(font_name)
    for installed in installed_fonts:
        normalized = normalize_font_name(installed)
        if target == normalized:
            return True
        if target and normalized and (target in normalized or normalized in target):
            return True
    return False


def detect_title_candidate(prs: Presentation, slide_index: int) -> dict[str, Any]:
    slide = prs.slides[slide_index - 1]
    candidates = []
    cover = slide_index == 1
    for shape_idx, shape in enumerate(slide.shapes, start=1):
        text = safe_text(shape)
        if not text:
            continue
        font_size = max_font_size_pt(shape)
        top_ratio = shape.top / prs.slide_height
        left_ratio = shape.left / prs.slide_width
        width_ratio = shape.width / prs.slide_width
        height_ratio = shape.height / prs.slide_height
        text_len = len(text.replace("\n", " ").strip())
        score = font_size
        if cover:
            score += font_size * 0.8
            score -= abs(0.42 - top_ratio) * 20
        else:
            score += max(0.0, 0.35 - top_ratio) * 65
            score += max(0.0, 0.30 - left_ratio) * 12
            if top_ratio < 0.18:
                score += 8
        if text_len > 90:
            score -= 18
        if width_ratio > 0.88 and text_len > 55:
            score -= 8
        if " / " in text or "\n" in text:
            score -= 2
        candidates.append(
            {
                "shape_index": shape_idx,
                "text": text,
                "score": round(score, 3),
                "top_ratio": round(top_ratio, 4),
                "left_ratio": round(left_ratio, 4),
                "width_ratio": round(width_ratio, 4),
                "height_ratio": round(height_ratio, 4),
                "font_size_pt": round(font_size, 2),
                "bbox": {
                    "left": int(shape.left),
                    "top": int(shape.top),
                    "width": int(shape.width),
                    "height": int(shape.height),
                },
            }
        )
    candidates.sort(key=lambda item: item["score"], reverse=True)
    best = candidates[0] if candidates else None
    second = candidates[1] if len(candidates) > 1 else None
    confidence = 0.0 if not best else 1.0 if not second else round(min(1.0, 0.55 + (best["score"] - second["score"]) / 24), 3)
    return {
        "slide_index": slide_index,
        "title": best["text"] if best else "",
        "shape_index": best["shape_index"] if best else None,
        "confidence": confidence,
        "candidates": candidates[:5],
    }


def infer_slide_role(slide_index: int, title_text: str, font_size_pt: float, text_shape_count: int) -> str:
    if slide_index == 1:
        return "cover"
    normalized = title_text.strip().upper()
    if normalized in {"Q&A", "THANK YOU", "THANK YOU!", "Q&A!"}:
        return "section"
    if text_shape_count <= 2 and font_size_pt >= 36:
        return "section"
    return "content"


def analyze_slide(prs: Presentation, slide_index: int, installed_fonts: set[str]) -> dict[str, Any]:
    slide = prs.slides[slide_index - 1]
    title_detection = detect_title_candidate(prs, slide_index)
    type_counts = Counter(getattr(shape.shape_type, "name", str(shape.shape_type)) for shape in slide.shapes)
    fonts = set()
    top_text_shapes = 0
    for shape in slide.shapes:
        text = safe_text(shape)
        if not text:
            continue
        if shape.top / prs.slide_height < 0.18:
            top_text_shapes += 1
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() and run.font.name:
                    fonts.add(run.font.name)

    role = infer_slide_role(
        slide_index,
        title_detection["title"],
        title_detection["candidates"][0]["font_size_pt"] if title_detection["candidates"] else 0.0,
        sum(1 for shape in slide.shapes if safe_text(shape)),
    )

    risks = []
    if not title_detection["title"]:
        risks.append("no_title_detected")
    elif title_detection["confidence"] < 0.78:
        risks.append("low_title_confidence")
    if type_counts["TABLE"] > 0:
        risks.append("table_present")
    if type_counts["CHART"] > 0:
        risks.append("chart_present")
    if type_counts["GROUP"] > 0:
        risks.append("group_shapes_present")
    if type_counts["PICTURE"] >= 3:
        risks.append("image_heavy")
    if sum(1 for shape in slide.shapes if safe_text(shape)) >= 12:
        risks.append("dense_text_layout")
    if top_text_shapes >= 3:
        risks.append("header_collision_risk")
    if len(title_detection["title"]) > 36:
        risks.append("long_title")

    missing_fonts = [font for font in sorted(fonts) if not font_is_available(font, installed_fonts)]
    if missing_fonts:
        risks.append("missing_fonts")

    return {
        "slide_index": slide_index,
        "title_detection": title_detection,
        "role": role,
        "shape_counts": dict(type_counts),
        "font_names": sorted(fonts),
        "missing_fonts": missing_fonts,
        "risks": sorted(set(risks)),
    }


def analyze_deck(source_path: Path, template_path: Path | None = None) -> dict[str, Any]:
    prs = Presentation(source_path)
    installed_fonts = get_installed_fonts()
    source_fonts = collect_fonts(prs)
    template_fonts = []
    if template_path and template_path.exists():
        template_fonts = collect_fonts(Presentation(template_path))

    slides = [analyze_slide(prs, idx, installed_fonts) for idx in range(1, len(prs.slides) + 1)]
    summary_counts = Counter()
    for slide in slides:
        for risk in slide["risks"]:
            summary_counts[risk] += 1

    missing_fonts = sorted({font for slide in slides for font in slide["missing_fonts"]})
    return {
        "source_path": str(source_path),
        "template_path": str(template_path) if template_path else "",
        "slide_count": len(prs.slides),
        "slide_size": {"width": int(prs.slide_width), "height": int(prs.slide_height)},
        "source_fonts": source_fonts,
        "template_fonts": template_fonts,
        "installed_font_count": len(installed_fonts),
        "missing_fonts": missing_fonts,
        "summary": {
            "risk_counts": dict(summary_counts),
            "high_risk_slides": [slide["slide_index"] for slide in slides if len(slide["risks"]) >= 2 or "low_title_confidence" in slide["risks"]],
        },
        "slides": slides,
    }
