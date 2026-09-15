from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation

from .analysis import max_font_size_pt, safe_text


@dataclass(frozen=True)
class HeaderStyle:
    left_ratio: float
    top_ratio: float
    width_ratio: float
    height_ratio: float
    margin_left: int
    margin_right: int
    margin_top: int
    margin_bottom: int
    font_name: str
    font_size_pt: float
    color_hex: str


@dataclass(frozen=True)
class TemplateAssets:
    cover_background_blob: bytes | None
    content_background_blob: bytes | None
    section_background_blob: bytes | None
    header_style_compact: HeaderStyle
    header_style_wide: HeaderStyle
    cover_layout_name: str
    toc_layout_name: str
    chapter_layout_name: str
    content_layout_name: str
    section_layout_name: str
    closing_layout_name: str


def _layout_picture_blob(layout) -> bytes | None:
    for shape in layout.shapes:
        if getattr(shape.shape_type, "name", str(shape.shape_type)) == "PICTURE":
            return shape.image.blob
    return None


def _slide_layout_name(slide) -> str:
    try:
        return slide.slide_layout.name or ""
    except Exception:
        return ""


def _shape_font_size_pt(shape) -> float:
    size = max_font_size_pt(shape)
    if size:
        return round(size, 2)

    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for xpath in (".//a:r/a:rPr", ".//a:defRPr"):
        node = shape._element.find(xpath, namespaces)
        if node is not None and node.get("sz"):
            return round(int(node.get("sz")) / 100, 2)
    return 32.0


def _shape_font_name(shape) -> str:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip() and run.font.name:
                return run.font.name

    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for xpath in (".//a:r/a:rPr/a:latin", ".//a:defRPr/a:latin"):
        node = shape._element.find(xpath, namespaces)
        if node is not None and node.get("typeface"):
            return node.get("typeface")
    return "腾讯体 W7"


def _shape_color_hex(shape) -> str:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            try:
                if run.font.color.rgb:
                    return str(run.font.color.rgb)
            except Exception:
                continue

    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for xpath in (".//a:r/a:rPr/a:solidFill/a:srgbClr", ".//a:defRPr/a:solidFill/a:srgbClr"):
        node = shape._element.find(xpath, namespaces)
        if node is not None and node.get("val"):
            return node.get("val")
    return "42D7FF"


def _header_candidate_kind(shape, prs: Presentation) -> str | None:
    text = safe_text(shape)
    if not text:
        return None
    left_ratio = shape.left / prs.slide_width
    top_ratio = shape.top / prs.slide_height
    width_ratio = shape.width / prs.slide_width
    height_ratio = shape.height / prs.slide_height
    if top_ratio >= 0.12 or left_ratio >= 0.16 or len(text) > 24:
        return None
    if not (0.03 <= height_ratio <= 0.12):
        return None
    if 0.08 <= width_ratio <= 0.35:
        return "compact"
    if 0.35 < width_ratio <= 0.7:
        return "wide"
    return None


def _build_header_style(shape, prs: Presentation) -> HeaderStyle:
    text_frame = shape.text_frame
    return HeaderStyle(
        left_ratio=round(shape.left / prs.slide_width, 4),
        top_ratio=round(shape.top / prs.slide_height, 4),
        width_ratio=round(shape.width / prs.slide_width, 4),
        height_ratio=round(shape.height / prs.slide_height, 4),
        margin_left=text_frame.margin_left,
        margin_right=text_frame.margin_right,
        margin_top=text_frame.margin_top,
        margin_bottom=text_frame.margin_bottom,
        font_name=_shape_font_name(shape),
        font_size_pt=_shape_font_size_pt(shape),
        color_hex=_shape_color_hex(shape),
    )


def _fallback_header_style() -> HeaderStyle:
    return HeaderStyle(
        left_ratio=0.0716,
        top_ratio=0.039,
        width_ratio=0.2275,
        height_ratio=0.0566,
        margin_left=45718,
        margin_right=45718,
        margin_top=45718,
        margin_bottom=45718,
        font_name="腾讯体 W7",
        font_size_pt=32.0,
        color_hex="42D7FF",
    )


def choose_template_layout_name(
    slide_analysis: dict[str, Any],
    assets: TemplateAssets,
) -> str:
    slide_index = slide_analysis["slide_index"]
    role = slide_analysis["role"]
    title_text = (slide_analysis.get("title_detection", {}) or {}).get("title", "")
    normalized_title = title_text.strip().upper().replace(" ", "")

    if role == "cover":
        return assets.cover_layout_name or assets.closing_layout_name
    if role == "section":
        if normalized_title in {"THANKS", "THANKYOU", "THANKYOU!", "Q&A", "Q&A!"}:
            return assets.closing_layout_name or assets.cover_layout_name
        return assets.section_layout_name or assets.chapter_layout_name or assets.content_layout_name
    if slide_index == 2 and assets.toc_layout_name:
        return assets.toc_layout_name
    if slide_index == 3 and assets.chapter_layout_name:
        return assets.chapter_layout_name
    return assets.content_layout_name or assets.chapter_layout_name or assets.toc_layout_name


def extract_template_assets(template_path: Path) -> TemplateAssets:
    prs = Presentation(template_path)
    slides = list(prs.slides)
    compact_style = None
    wide_style = None

    for slide in slides[2:]:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            kind = _header_candidate_kind(shape, prs)
            if kind == "compact" and compact_style is None:
                compact_style = _build_header_style(shape, prs)
            elif kind == "wide" and wide_style is None:
                wide_style = _build_header_style(shape, prs)
        if compact_style and wide_style:
            break

    compact_style = compact_style or _fallback_header_style()
    wide_style = wide_style or HeaderStyle(
        left_ratio=compact_style.left_ratio,
        top_ratio=compact_style.top_ratio,
        width_ratio=max(compact_style.width_ratio, 0.58),
        height_ratio=max(compact_style.height_ratio, 0.1045),
        margin_left=compact_style.margin_left,
        margin_right=compact_style.margin_right,
        margin_top=compact_style.margin_top,
        margin_bottom=compact_style.margin_bottom,
        font_name=compact_style.font_name,
        font_size_pt=compact_style.font_size_pt,
        color_hex=compact_style.color_hex,
    )

    layout_names = [_slide_layout_name(slide) for slide in slides]
    layout_counter = Counter(name for name in layout_names[3:-1] if name)
    content_layout_name = layout_counter.most_common(1)[0][0] if layout_counter else next((name for name in layout_names if name), "")
    toc_layout_name = layout_names[1] if len(layout_names) >= 2 else content_layout_name
    chapter_layout_name = layout_names[2] if len(layout_names) >= 3 else content_layout_name
    cover_layout_name = layout_names[0] if layout_names else content_layout_name
    closing_layout_name = layout_names[-1] if layout_names else cover_layout_name
    section_layout_name = next((name for name in layout_names if "章节" in name), chapter_layout_name or content_layout_name)

    cover_background = _layout_picture_blob(slides[1].slide_layout) if len(slides) >= 2 else None
    content_slide = next((slide for slide in slides[2:] if slide.slide_layout is not None), None)
    content_background = _layout_picture_blob(content_slide.slide_layout) if content_slide else None
    section_slide = None
    for slide in slides:
        first_shape = slide.shapes[0] if len(slide.shapes) else None
        if first_shape and safe_text(first_shape).strip().upper() in {"Q&A", "THANK YOU!"}:
            section_slide = slide
            break
    section_background = _layout_picture_blob(section_slide.slide_layout) if section_slide else content_background

    return TemplateAssets(
        cover_background_blob=cover_background,
        content_background_blob=content_background,
        section_background_blob=section_background,
        header_style_compact=compact_style,
        header_style_wide=wide_style,
        cover_layout_name=cover_layout_name,
        toc_layout_name=toc_layout_name,
        chapter_layout_name=chapter_layout_name,
        content_layout_name=content_layout_name,
        section_layout_name=section_layout_name,
        closing_layout_name=closing_layout_name,
    )
