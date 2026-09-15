from __future__ import annotations

import json
import statistics
from pathlib import Path

from pptx import Presentation

from .analysis import max_font_size_pt, safe_text
from .config import TemplateProfile
from .template_assets import extract_template_assets


def extract_template_profile(template_path: Path) -> TemplateProfile:
    try:
        assets = extract_template_assets(template_path)
        safe_zone_bottom = max(
            assets.header_style_compact.top_ratio + assets.header_style_compact.height_ratio,
            assets.header_style_wide.top_ratio + assets.header_style_wide.height_ratio,
        ) + 0.04
        return TemplateProfile(
            header_left_ratio=round(assets.header_style_compact.left_ratio, 4),
            header_top_ratio=round(assets.header_style_compact.top_ratio, 4),
            header_width_ratio=round(assets.header_style_compact.width_ratio, 4),
            header_height_ratio=round(assets.header_style_compact.height_ratio, 4),
            accent_left_ratio=0.022,
            accent_top_ratio=0.041,
            accent_width_ratio=0.018,
            accent_height_ratio=0.014,
            safe_zone_bottom_ratio=round(min(0.22, max(0.12, safe_zone_bottom)), 4),
        )
    except Exception:
        pass

    prs = Presentation(template_path)
    lefts = []
    tops = []
    widths = []
    heights = []
    safe_zone_candidates = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        if slide_index <= 2:
            continue
        text_shapes = []
        for shape in slide.shapes:
            text = safe_text(shape)
            if not text:
                continue
            text_shapes.append((shape, text, max_font_size_pt(shape)))
        for shape, text, font_size in text_shapes:
            left_ratio = shape.left / prs.slide_width
            top_ratio = shape.top / prs.slide_height
            width_ratio = shape.width / prs.slide_width
            height_ratio = shape.height / prs.slide_height
            if (
                top_ratio < 0.12
                and left_ratio < 0.16
                and 10 <= font_size <= 36
                and len(text) <= 24
                and 0.08 <= width_ratio <= 0.35
                and 0.03 <= height_ratio <= 0.09
            ):
                lefts.append(left_ratio)
                tops.append(top_ratio)
                widths.append(width_ratio)
                heights.append(height_ratio)
                safe_zone_candidates.append(top_ratio + height_ratio + 0.04)

    if not lefts:
        return TemplateProfile()

    return TemplateProfile(
        header_left_ratio=round(statistics.median(lefts), 4),
        header_top_ratio=round(statistics.median(tops), 4),
        header_width_ratio=round(statistics.median(widths), 4),
        header_height_ratio=round(statistics.median(heights), 4),
        accent_left_ratio=round(max(0.015, statistics.median(lefts) * 0.31), 4),
        accent_top_ratio=round(max(0.02, statistics.median(tops) * 1.1), 4),
        accent_width_ratio=0.018,
        accent_height_ratio=0.014,
        safe_zone_bottom_ratio=round(min(0.22, max(0.12, statistics.median(safe_zone_candidates))), 4),
    )


def write_template_profile_json(output_dir: Path, profile: TemplateProfile) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "template_profile.json"
    path.write_text(json.dumps(profile.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
