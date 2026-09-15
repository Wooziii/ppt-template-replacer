from __future__ import annotations

import colorsys
import json
from collections import Counter
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from PIL import Image
from pptx import Presentation

from .config import TemplatePalette
from .template_assets import TemplateAssets, extract_template_assets


@dataclass(frozen=True)
class TemplateTheme:
    palette: TemplatePalette
    metadata: dict[str, Any]


def hex_to_rgb_tuple(hex_color: str) -> tuple[int, int, int]:
    value = hex_color.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in range(0, 6, 2))


def rgb_tuple_to_hex(color: tuple[int, int, int]) -> str:
    return "".join(f"{max(0, min(255, channel)):02X}" for channel in color)


def luminance(color: tuple[int, int, int]) -> float:
    red, green, blue = [channel / 255.0 for channel in color]
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def mix(color_a: tuple[int, int, int], color_b: tuple[int, int, int], ratio: float) -> tuple[int, int, int]:
    ratio = max(0.0, min(1.0, ratio))
    return tuple(int(round(channel_a * (1.0 - ratio) + channel_b * ratio)) for channel_a, channel_b in zip(color_a, color_b))


def shift_value(color: tuple[int, int, int], delta: float) -> tuple[int, int, int]:
    red, green, blue = [channel / 255.0 for channel in color]
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    value = max(0.0, min(1.0, value + delta))
    rgb_float = colorsys.hsv_to_rgb(hue, saturation, value)
    return tuple(int(round(channel * 255.0)) for channel in rgb_float)


def tune_hsv(
    color: tuple[int, int, int],
    *,
    saturation_scale: float = 1.0,
    value_scale: float = 1.0,
    value_delta: float = 0.0,
) -> tuple[int, int, int]:
    red, green, blue = [channel / 255.0 for channel in color]
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    saturation = max(0.0, min(1.0, saturation * saturation_scale))
    value = max(0.0, min(1.0, value * value_scale + value_delta))
    rgb_float = colorsys.hsv_to_rgb(hue, saturation, value)
    return tuple(int(round(channel * 255.0)) for channel in rgb_float)


def flatten_shapes(shapes: Iterable) -> list:
    flat = []
    for shape in shapes:
        flat.append(shape)
        if getattr(shape, "shape_type", None) and getattr(shape.shape_type, "name", str(shape.shape_type)) == "GROUP":
            flat.extend(flatten_shapes(shape.shapes))
    return flat


def shape_fill_hex(shape) -> str | None:
    try:
        if shape.fill.type == 1 and shape.fill.fore_color.rgb:
            return str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def shape_line_hex(shape) -> str | None:
    try:
        if shape.line.color.rgb:
            return str(shape.line.color.rgb)
    except Exception:
        pass
    return None


def iter_text_colors(shape) -> list[str]:
    colors = []
    if not getattr(shape, "has_text_frame", False):
        return colors
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            try:
                if run.font.color.rgb:
                    colors.append(str(run.font.color.rgb))
            except Exception:
                continue
    return colors


def dominant_colors_from_blob(blob: bytes | None, *, colors: int = 8) -> list[str]:
    if not blob:
        return []
    image = Image.open(BytesIO(blob)).convert("RGB")
    image.thumbnail((240, 135))
    palette = image.quantize(colors=colors, method=Image.MEDIANCUT).convert("RGB")
    raw = palette.getcolors()
    if not raw:
        return []
    return [rgb_tuple_to_hex(color) for _, color in sorted(raw, reverse=True)]


def pick_background_hex(background_candidates: list[str], fill_counter: Counter[str]) -> str:
    if background_candidates:
        return background_candidates[0]
    if fill_counter:
        return fill_counter.most_common(1)[0][0]
    return "0B0A18"


def pick_text_hex(
    counter: Counter[str],
    *,
    minimum_luma: float | None = None,
    maximum_luma: float | None = None,
    fallback: str,
) -> str:
    for hex_color, _ in counter.most_common():
        luma = luminance(hex_to_rgb_tuple(hex_color))
        if minimum_luma is not None and luma < minimum_luma:
            continue
        if maximum_luma is not None and luma > maximum_luma:
            continue
        return hex_color
    return fallback


def pick_accent_hex(title_hex: str, line_counter: Counter[str], fill_counter: Counter[str], text_counter: Counter[str]) -> str:
    for counter in (fill_counter, line_counter, text_counter):
        for hex_color, _ in counter.most_common():
            red, green, blue = hex_to_rgb_tuple(hex_color)
            hue, saturation, _ = colorsys.rgb_to_hsv(red / 255.0, green / 255.0, blue / 255.0)
            if saturation < 0.35:
                continue
            if 0.06 <= hue <= 0.18:
                return hex_color
    return "FF4A2D"


def pick_blue_hex(background_hex: str, line_counter: Counter[str]) -> str:
    for hex_color, _ in line_counter.most_common():
        red, green, blue = hex_to_rgb_tuple(hex_color)
        hue, saturation, _ = colorsys.rgb_to_hsv(red / 255.0, green / 255.0, blue / 255.0)
        if saturation >= 0.2 and 0.48 <= hue <= 0.68:
            return hex_color
    return rgb_tuple_to_hex(shift_value(hex_to_rgb_tuple(background_hex), 0.12))


def extract_template_theme(template_path: Path, assets: TemplateAssets | None = None) -> TemplateTheme:
    assets = assets or extract_template_assets(template_path)
    prs = Presentation(template_path)
    slides = list(prs.slides)
    candidate_slides = slides[2:] or slides

    fill_counter: Counter[str] = Counter()
    line_counter: Counter[str] = Counter()
    text_counter: Counter[str] = Counter()

    for slide in candidate_slides:
        for shape in flatten_shapes(slide.shapes):
            fill_hex = shape_fill_hex(shape)
            if fill_hex:
                fill_counter[fill_hex] += 1
            line_hex = shape_line_hex(shape)
            if line_hex:
                line_counter[line_hex] += 1
            for text_hex in iter_text_colors(shape):
                text_counter[text_hex] += 1

    background_candidates = dominant_colors_from_blob(assets.content_background_blob)
    background_hex = pick_background_hex(background_candidates, fill_counter)
    background_rgb = hex_to_rgb_tuple(background_hex)
    is_dark_theme = luminance(background_rgb) < 0.45

    title_hex = assets.header_style_compact.color_hex or "42D7FF"
    title_rgb = hex_to_rgb_tuple(title_hex)
    accent_hex = pick_accent_hex(title_hex, line_counter, fill_counter, text_counter)
    accent_rgb = hex_to_rgb_tuple(accent_hex)

    text_light = pick_text_hex(text_counter, minimum_luma=0.82, fallback="F6FAFF")
    text_dark = pick_text_hex(text_counter, maximum_luma=0.36, fallback="1E293B")
    text_dark_rgb = hex_to_rgb_tuple(text_dark)
    softened_title_rgb = tune_hsv(title_rgb, saturation_scale=0.72, value_scale=0.9)
    softened_accent_rgb = tune_hsv(accent_rgb, saturation_scale=0.7, value_scale=0.92)

    if is_dark_theme:
        panel_dark = rgb_tuple_to_hex(mix(background_rgb, (255, 255, 255), 0.10))
        panel_mid = rgb_tuple_to_hex(mix(background_rgb, (255, 255, 255), 0.16))
        panel_alt = rgb_tuple_to_hex(mix(background_rgb, title_rgb, 0.14))
        panel_teal = rgb_tuple_to_hex(mix(hex_to_rgb_tuple(panel_mid), softened_title_rgb, 0.18))
        panel_warm = rgb_tuple_to_hex(mix(hex_to_rgb_tuple(panel_mid), softened_accent_rgb, 0.22))
        border_muted = rgb_tuple_to_hex(mix(background_rgb, softened_title_rgb, 0.42))
    else:
        surface_rgb = mix(background_rgb, (255, 255, 255), 0.68)
        panel_mid = rgb_tuple_to_hex(surface_rgb)
        panel_dark = rgb_tuple_to_hex(mix(surface_rgb, text_dark_rgb, 0.06))
        panel_alt = rgb_tuple_to_hex(mix(surface_rgb, softened_title_rgb, 0.08))
        panel_teal = rgb_tuple_to_hex(mix(surface_rgb, softened_title_rgb, 0.12))
        panel_warm = rgb_tuple_to_hex(mix(surface_rgb, softened_accent_rgb, 0.12))
        border_seed_rgb = mix(text_dark_rgb, softened_title_rgb, 0.30)
        border_muted = rgb_tuple_to_hex(mix(surface_rgb, border_seed_rgb, 0.55))

    chart_teal = title_hex
    chart_gold = accent_hex
    chart_blue = pick_blue_hex(background_hex, line_counter)
    chart_series = (chart_teal, chart_gold, chart_blue, title_hex, border_muted)

    palette = TemplatePalette(
        background_dark=background_hex,
        title_cyan=title_hex,
        accent_orange=accent_hex,
        text_light=text_light,
        text_dark=text_dark,
        text_muted=rgb_tuple_to_hex(mix(text_dark_rgb, background_rgb, 0.32 if is_dark_theme else 0.46)),
        panel_dark=panel_dark,
        panel_mid=panel_mid,
        panel_alt=panel_alt,
        panel_teal=panel_teal,
        panel_warm=panel_warm,
        border_muted=border_muted,
        chart_teal=chart_teal,
        chart_gold=chart_gold,
        chart_blue=chart_blue,
        body_font=TemplatePalette().body_font,
        heading_font=assets.header_style_compact.font_name or TemplatePalette().heading_font,
        chart_series=chart_series,
    )

    metadata = {
        "template_path": str(template_path),
        "is_dark_theme": is_dark_theme,
        "background_candidates": background_candidates[:6],
        "sampled_fill_colors": [color for color, _ in fill_counter.most_common(8)],
        "sampled_line_colors": [color for color, _ in line_counter.most_common(8)],
        "sampled_text_colors": [color for color, _ in text_counter.most_common(8)],
        "header_color": title_hex,
        "background_color": background_hex,
        "accent_color": accent_hex,
        "surface_mode": "dark" if is_dark_theme else "light",
    }
    return TemplateTheme(palette=palette, metadata=metadata)


def write_template_palette_json(output_dir: Path, theme: TemplateTheme) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "template_palette.json"
    payload = {
        "palette": asdict(theme.palette),
        "metadata": theme.metadata,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
