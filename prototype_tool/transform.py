from __future__ import annotations

import colorsys
import math
import re
import shutil
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .analysis import safe_text
from .config import TemplatePalette, TemplateProfile
from .template_assets import HeaderStyle, TemplateAssets, extract_template_assets


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def hex_to_rgb_tuple(hex_color: str) -> tuple[int, int, int]:
    value = hex_color.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in range(0, 6, 2))


def luminance_from_rgb_tuple(color: tuple[int, int, int]) -> float:
    red, green, blue = [channel / 255.0 for channel in color]
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def ppt_rgb_to_tuple(color: RGBColor | None) -> tuple[int, int, int] | None:
    if color is None:
        return None
    value = str(color)
    if len(value) != 6:
        return None
    return tuple(int(value[index : index + 2], 16) for index in range(0, 6, 2))


def intersection_ratio(shape_a, shape_b) -> float:
    left = max(shape_a.left, shape_b.left)
    top = max(shape_a.top, shape_b.top)
    right = min(shape_a.left + shape_a.width, shape_b.left + shape_b.width)
    bottom = min(shape_a.top + shape_a.height, shape_b.top + shape_b.height)
    if right <= left or bottom <= top:
        return 0.0
    inter = (right - left) * (bottom - top)
    area = max(1, shape_a.width * shape_a.height)
    return inter / area


def shape_local_bg_is_light(shape, slide_bg_luma: float, underlying_shapes: list) -> bool:
    try:
        fill = shape.fill
        if fill.type == 1 and fill.fore_color.rgb:
            color = ppt_rgb_to_tuple(fill.fore_color.rgb)
            if color:
                return luminance_from_rgb_tuple(color) > 0.75
    except Exception:
        pass
    for under in reversed(underlying_shapes):
        try:
            if intersection_ratio(shape, under) < 0.45:
                continue
            if under.fill.type == 1 and under.fill.fore_color.rgb:
                color = ppt_rgb_to_tuple(under.fill.fore_color.rgb)
                if color:
                    return luminance_from_rgb_tuple(color) > 0.75
        except Exception:
            continue
    return slide_bg_luma > 0.55


def recolor_text_shape(
    shape,
    target_hex: str,
    *,
    font_name: str | None = None,
    bold: bool | None = None,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    recolor_text_frame(shape.text_frame, target_hex, font_name=font_name, bold=bold)


def _set_typeface_node(parent, tag_name: str, font_name: str) -> None:
    node = parent.find(qn(tag_name))
    if node is None:
        node = OxmlElement(tag_name)
        parent.append(node)
    node.set("typeface", font_name)


def apply_font_family_to_run(run, font_name: str) -> None:
    run.font.name = font_name
    r_pr = run._r.get_or_add_rPr()
    _set_typeface_node(r_pr, "a:latin", font_name)
    _set_typeface_node(r_pr, "a:ea", font_name)
    _set_typeface_node(r_pr, "a:cs", font_name)


def apply_font_family_to_paragraph(paragraph, font_name: str) -> None:
    try:
        paragraph.font.name = font_name
    except Exception:
        pass
    for run in paragraph.runs:
        apply_font_family_to_run(run, font_name)


def recolor_text_frame(text_frame, target_hex: str, font_name: str | None = None, bold: bool | None = None) -> None:
    for paragraph in text_frame.paragraphs:
        if not paragraph.runs:
            paragraph.font.color.rgb = rgb(target_hex)
            if font_name:
                apply_font_family_to_paragraph(paragraph, font_name)
            if bold is not None:
                paragraph.font.bold = bold
        for run in paragraph.runs:
            run.font.color.rgb = rgb(target_hex)
            if font_name:
                apply_font_family_to_run(run, font_name)
            if bold is not None:
                run.font.bold = bold


def set_shape_fill_color(shape, fill_hex: str) -> None:
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill_hex)
    except Exception:
        pass


def set_shape_line_color(shape, line_hex: str) -> None:
    try:
        shape.line.color.rgb = rgb(line_hex)
    except Exception:
        try:
            shape.line.fill.solid()
            shape.line.fill.fore_color.rgb = rgb(line_hex)
        except Exception:
            pass


def get_shape_solid_fill_tuple(shape) -> tuple[int, int, int] | None:
    try:
        if shape.fill.type == 1 and shape.fill.fore_color.rgb:
            return ppt_rgb_to_tuple(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def get_shape_line_tuple(shape) -> tuple[int, int, int] | None:
    try:
        if shape.line.color.rgb:
            return ppt_rgb_to_tuple(shape.line.color.rgb)
    except Exception:
        pass
    return None


def shape_area_ratio(shape, slide_w: int, slide_h: int) -> float:
    return (shape.width * shape.height) / max(1, slide_w * slide_h)


def classify_color(color: tuple[int, int, int]) -> tuple[float, float, float]:
    red, green, blue = [channel / 255.0 for channel in color]
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    return hue, saturation, value


def palette_background_is_light(palette: TemplatePalette) -> bool:
    return luminance_from_rgb_tuple(hex_to_rgb_tuple(palette.background_dark)) > 0.55


def choose_text_hex_for_fill(fill_hex: str, palette: TemplatePalette) -> str:
    fill_luma = luminance_from_rgb_tuple(hex_to_rgb_tuple(fill_hex))
    return palette.text_dark if fill_luma >= 0.56 else palette.text_light


def choose_text_hex_for_shape(
    shape,
    palette: TemplatePalette,
    slide_bg_luma: float,
    underlying_shapes: list,
    *,
    fill_hex: str | None = None,
) -> str:
    if fill_hex:
        return choose_text_hex_for_fill(fill_hex, palette)
    return palette.text_dark if shape_local_bg_is_light(shape, slide_bg_luma, underlying_shapes) else palette.text_light


def map_fill_color(source_color: tuple[int, int, int] | None, palette: TemplatePalette, *, large_shape: bool) -> str:
    if source_color is None:
        return palette.panel_mid if large_shape else palette.panel_alt
    hue, saturation, value = classify_color(source_color)
    luma = luminance_from_rgb_tuple(source_color)
    if saturation < 0.12:
        if large_shape:
            return palette.panel_mid if luma > 0.42 else palette.panel_dark
        return palette.border_muted if luma > 0.6 else palette.panel_alt
    if 0.07 <= hue <= 0.17:
        return palette.panel_warm if large_shape else palette.chart_gold
    if 0.18 < hue <= 0.46:
        return palette.panel_teal if large_shape else palette.chart_teal
    if 0.46 < hue <= 0.74:
        return palette.panel_alt if large_shape else palette.title_cyan
    return palette.panel_alt if large_shape else palette.chart_blue


def map_line_color(source_color: tuple[int, int, int] | None, palette: TemplatePalette) -> str:
    if source_color is None:
        return palette.border_muted
    hue, saturation, _ = classify_color(source_color)
    if saturation < 0.12:
        return palette.border_muted
    if 0.07 <= hue <= 0.17:
        return palette.chart_gold
    if 0.18 < hue <= 0.46:
        return palette.chart_teal
    if 0.46 < hue <= 0.74:
        return palette.title_cyan
    return palette.chart_blue


def has_explicit_fill(shape) -> bool:
    try:
        return shape.fill.type == 1 and bool(shape.fill.fore_color.rgb)
    except Exception:
        return False


def has_visible_line(shape) -> bool:
    try:
        return bool(shape.line.color.rgb)
    except Exception:
        try:
            return bool(shape.line.fill.fore_color.rgb)
        except Exception:
            return False


def estimate_text_units(text: str) -> float:
    units = 0.0
    for character in text:
        if ord(character) > 127:
            units += 1.0
        elif character.isupper():
            units += 0.72
        elif character.islower():
            units += 0.58
        elif character.isdigit():
            units += 0.6
        elif character in {" ", "/", "-", "_", ".", ",", ":", ";", "|"}:
            units += 0.32
        else:
            units += 0.5
    return units


def clear_shape_fill_and_line(shape) -> None:
    try:
        shape.fill.background()
    except Exception:
        pass
    try:
        shape.line.fill.background()
    except Exception:
        pass


def should_enable_text_autofit(shape, slide_w: int, slide_h: int) -> bool:
    text = normalize_header_text(safe_text(shape))
    if not text or len(text) > 180:
        return False
    if text.count("\n") > 6:
        return False
    area_ratio = shape_area_ratio(shape, slide_w, slide_h)
    return area_ratio < 0.11 or shape.height < slide_h * 0.14 or shape.width < slide_w * 0.42


def apply_text_autofit(shape, slide_w: int, slide_h: int) -> None:
    if not should_enable_text_autofit(shape, slide_w, slide_h):
        return
    try:
        shape.text_frame.word_wrap = True
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


def is_nested_text_panel(shape, underlying_shapes: list) -> bool:
    if not getattr(shape, "has_text_frame", False) or not safe_text(shape):
        return False
    shape_area = max(1, shape.width * shape.height)
    for under in reversed(underlying_shapes):
        if getattr(under, "shape_type", None) and getattr(under.shape_type, "name", str(under.shape_type)) == "PICTURE":
            continue
        if not has_explicit_fill(under):
            continue
        under_area = max(1, under.width * under.height)
        if under_area < shape_area * 1.35:
            continue
        overlap_ratio = intersection_ratio(shape, under)
        if overlap_ratio < 0.84:
            banner_overlap = overlap_ratio >= 0.45 and shape.height < under.height * 0.25
            if not banner_overlap:
                continue
        return True
    return False


def ensure_cell_borders(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        line = tc_pr.find(qn(edge))
        if line is None:
            line = OxmlElement(edge)
            tc_pr.append(line)
        for child in list(line):
            line.remove(child)
        line.set("w", "12700")
        line.set("cap", "flat")
        line.set("cmpd", "sng")
        line.set("algn", "ctr")
        solid_fill = OxmlElement("a:solidFill")
        srgb_clr = OxmlElement("a:srgbClr")
        srgb_clr.set("val", color_hex)
        solid_fill.append(srgb_clr)
        line.append(solid_fill)
        prst_dash = OxmlElement("a:prstDash")
        prst_dash.set("val", "solid")
        line.append(prst_dash)
        line.append(OxmlElement("a:round"))
        head_end = OxmlElement("a:headEnd")
        head_end.set("type", "none")
        head_end.set("w", "med")
        head_end.set("len", "med")
        line.append(head_end)
        tail_end = OxmlElement("a:tailEnd")
        tail_end.set("type", "none")
        tail_end.set("w", "med")
        tail_end.set("len", "med")
        line.append(tail_end)


def style_table_shape(shape, palette: TemplatePalette) -> dict[str, Any]:
    table = shape.table
    row_count = len(table.rows)
    col_count = len(table.columns)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            is_header_row = row_idx == 0 and row_count > 1
            is_header_col = col_idx == 0 and col_count > 2 and row_idx > 0
            if is_header_row:
                fill_hex = palette.panel_alt
                text_hex = choose_text_hex_for_fill(fill_hex, palette)
                font_name = palette.heading_font
                bold = True
            elif is_header_col:
                fill_hex = palette.panel_mid
                text_hex = choose_text_hex_for_fill(fill_hex, palette)
                font_name = palette.heading_font
                bold = True
            else:
                fill_hex = palette.panel_dark if row_idx % 2 else palette.panel_mid
                text_hex = choose_text_hex_for_fill(fill_hex, palette)
                font_name = palette.body_font
                bold = False
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill_hex)
            ensure_cell_borders(cell, palette.border_muted)
            recolor_text_frame(cell.text_frame, text_hex, font_name=font_name, bold=bold)
    return {"rows": row_count, "cols": col_count}


def style_chart_text(text_frame, color_hex: str, font_name: str) -> None:
    for paragraph in text_frame.paragraphs:
        if not paragraph.runs:
            paragraph.font.color.rgb = rgb(color_hex)
            apply_font_family_to_paragraph(paragraph, font_name)
        for run in paragraph.runs:
            run.font.color.rgb = rgb(color_hex)
            apply_font_family_to_run(run, font_name)


def style_chart_axis(axis, palette: TemplatePalette, text_hex: str) -> None:
    try:
        axis.format.line.color.rgb = rgb(palette.border_muted)
    except Exception:
        pass
    try:
        axis.tick_labels.font.color.rgb = rgb(text_hex)
        axis.tick_labels.font.name = palette.body_font
    except Exception:
        pass
    try:
        axis.has_major_gridlines
        axis.major_gridlines.format.line.color.rgb = rgb(palette.border_muted)
    except Exception:
        pass
    try:
        if axis.has_title:
            style_chart_text(axis.axis_title.text_frame, text_hex, palette.heading_font)
    except Exception:
        pass


def style_chart_shape(shape, palette: TemplatePalette) -> dict[str, Any]:
    chart = shape.chart
    chart_text_hex = palette.text_dark if palette_background_is_light(palette) else palette.text_light
    try:
        chart.chart_area.format.fill.background()
    except Exception:
        pass
    try:
        chart.chart_area.format.line.color.rgb = rgb(palette.border_muted)
    except Exception:
        pass
    try:
        chart.plot_area.format.fill.background()
    except Exception:
        pass
    try:
        chart.plot_area.format.line.color.rgb = rgb(palette.border_muted)
    except Exception:
        pass
    try:
        if chart.has_title:
            style_chart_text(chart.chart_title.text_frame, chart_text_hex, palette.heading_font)
    except Exception:
        pass
    try:
        if chart.has_legend:
            chart.legend.font.color.rgb = rgb(chart_text_hex)
            chart.legend.font.name = palette.body_font
    except Exception:
        pass

    for axis_name in ("category_axis", "value_axis", "series_axis"):
        try:
            style_chart_axis(getattr(chart, axis_name), palette, chart_text_hex)
        except Exception:
            continue

    plot = chart.plots[0] if chart.plots else None
    if plot is not None:
        try:
            plot.has_data_labels = True
            plot.data_labels.font.color.rgb = rgb(chart_text_hex)
            plot.data_labels.font.name = palette.body_font
        except Exception:
            pass

    for idx, series in enumerate(chart.series):
        color_hex = palette.chart_series[idx % len(palette.chart_series)]
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(color_hex)
        except Exception:
            pass
        try:
            series.format.line.color.rgb = rgb(color_hex)
        except Exception:
            pass
        try:
            marker = series.marker
            marker.format.fill.solid()
            marker.format.fill.fore_color.rgb = rgb(color_hex)
            marker.format.line.color.rgb = rgb(color_hex)
        except Exception:
            pass
        if len(chart.series) == 1:
            for point_idx, point in enumerate(series.points):
                point_hex = palette.chart_blue
                if point_idx == len(series.points) - 1:
                    point_hex = palette.accent_orange
                try:
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = rgb(point_hex)
                except Exception:
                    pass
                try:
                    point.format.line.color.rgb = rgb(point_hex)
                except Exception:
                    pass
        try:
            if plot is not None and plot.has_data_labels:
                plot.data_labels.font.color.rgb = rgb(chart_text_hex)
        except Exception:
            pass
    return {"series_count": len(chart.series)}


def style_vector_shape(
    shape,
    palette: TemplatePalette,
    slide_w: int,
    slide_h: int,
    slide_bg_luma: float,
    underlying_shapes: list,
) -> dict[str, Any] | None:
    shape_name = getattr(shape.shape_type, "name", str(shape.shape_type))
    if shape_name not in {"AUTO_SHAPE", "FREEFORM", "LINE"}:
        return None

    has_text = getattr(shape, "has_text_frame", False) and bool(safe_text(shape))
    explicit_fill = has_explicit_fill(shape)
    visible_line = has_visible_line(shape)
    fill_tuple = get_shape_solid_fill_tuple(shape)
    line_tuple = get_shape_line_tuple(shape)
    large_shape = shape_name != "LINE" and (shape_area_ratio(shape, slide_w, slide_h) > 0.012 or has_text)
    transparent_text_container = has_text and shape_name != "LINE" and not explicit_fill

    if transparent_text_container:
        clear_shape_fill_and_line(shape)
        fill_hex = None
    elif shape_name != "LINE":
        fill_hex = map_fill_color(fill_tuple, palette, large_shape=large_shape)
        if explicit_fill or not has_text:
            set_shape_fill_color(shape, fill_hex)
    else:
        fill_hex = None

    if shape_name == "LINE" or visible_line or explicit_fill or not has_text:
        line_hex = map_line_color(line_tuple or fill_tuple, palette)
        set_shape_line_color(shape, line_hex)
    else:
        line_hex = None

    if has_text:
        text_hex = choose_text_hex_for_shape(
            shape,
            palette,
            slide_bg_luma,
            underlying_shapes,
            fill_hex=fill_hex,
        )
        font_name = palette.heading_font if large_shape else palette.body_font
        recolor_text_frame(shape.text_frame, text_hex, font_name=font_name)
        apply_text_autofit(shape, slide_w, slide_h)

    return {"shape_type": shape_name, "large_shape": large_shape, "transparent_text_container": transparent_text_container}


def style_group_shape(
    group_shape,
    palette: TemplatePalette,
    slide_w: int,
    slide_h: int,
    slide_bg_luma: float,
) -> dict[str, Any]:
    styled = {"vectors": 0, "tables": 0, "charts": 0, "text_nodes": 0}
    children = list(group_shape.shapes)
    for child_idx, child in enumerate(children, start=1):
        if getattr(child, "shape_type", None) and getattr(child.shape_type, "name", str(child.shape_type)) == "GROUP":
            nested = style_group_shape(child, palette, slide_w, slide_h, slide_bg_luma)
            for key, value in nested.items():
                styled[key] += value
            continue
        if getattr(child, "has_text_frame", False) and safe_text(child):
            if not has_explicit_fill(child) or is_nested_text_panel(child, children[: child_idx - 1]):
                clear_shape_fill_and_line(child)
        if getattr(child, "has_table", False):
            style_table_shape(child, palette)
            styled["tables"] += 1
            continue
        if getattr(child, "has_chart", False):
            style_chart_shape(child, palette)
            styled["charts"] += 1
            continue
        vector_meta = style_vector_shape(
            child,
            palette,
            slide_w,
            slide_h,
            slide_bg_luma,
            children[: child_idx - 1],
        )
        if vector_meta:
            styled["vectors"] += 1
        if getattr(child, "has_text_frame", False) and safe_text(child) and vector_meta is None:
            recolor_text_shape(
                child,
                choose_text_hex_for_shape(
                    child,
                    palette,
                    slide_bg_luma,
                    children[: child_idx - 1],
                ),
                font_name=palette.body_font,
            )
            apply_text_autofit(child, slide_w, slide_h)
            styled["text_nodes"] += 1
    return styled


def remove_shape(shape) -> None:
    sp = shape._element
    sp.getparent().remove(sp)


def shift_shape_down(shape, delta_emu: int, slide_h: int) -> None:
    new_top = min(int(shape.top + delta_emu), max(0, slide_h - shape.height))
    shape.top = new_top


def normalize_header_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def apply_template_background(slide, slide_w: int, slide_h: int, background_blob: bytes | None) -> bool:
    # python-pptx cannot attach a real PowerPoint master/layout to existing slides.
    # Avoid rasterizing the template background into a full-slide image so the deck
    # remains editable and does not masquerade as a native-template presentation.
    return False


def choose_header_style(title_text: str, assets: TemplateAssets) -> HeaderStyle:
    normalized = normalize_header_text(title_text)
    if len(normalized) > 18 or any(character.isascii() for character in normalized):
        return assets.header_style_wide
    return assets.header_style_compact


def choose_header_width_ratio(style: HeaderStyle, title_text: str) -> float:
    width_ratio = style.width_ratio
    if width_ratio < 0.35:
        return width_ratio
    normalized_text = normalize_header_text(title_text)
    text_units = estimate_text_units(normalized_text)
    ascii_chars = sum(1 for character in normalized_text if character.isascii() and not character.isspace())
    visible_chars = sum(1 for character in normalized_text if not character.isspace())
    ascii_ratio = ascii_chars / max(1, visible_chars)
    if ascii_ratio > 0.35 and text_units > 14:
        width_ratio += 0.05 + min(0.04, max(0.0, text_units - 14) * 0.003)
    elif text_units > 18:
        width_ratio += 0.04
    return round(min(0.7, width_ratio), 4)


def choose_fitted_header_font_size(
    style: HeaderStyle,
    slide_w: int,
    title_text: str,
    width_ratio: float | None = None,
) -> float:
    font_size = style.font_size_pt
    normalized_text = normalize_header_text(title_text)
    resolved_width_ratio = width_ratio if width_ratio is not None else choose_header_width_ratio(style, title_text)
    box_width_pt = ((slide_w * resolved_width_ratio) - style.margin_left - style.margin_right) / 12700
    text_units = estimate_text_units(title_text)
    ascii_chars = sum(1 for character in normalized_text if character.isascii() and not character.isspace())
    visible_chars = sum(1 for character in normalized_text if not character.isspace())
    ascii_ratio = ascii_chars / max(1, visible_chars)
    if ascii_ratio > 0.75:
        if visible_chars > 28:
            font_size = min(font_size, 24.0)
        elif visible_chars > 22:
            font_size = min(font_size, 27.0)
    if ascii_ratio > 0.75:
        width_factor = 0.78
    elif ascii_ratio > 0.35:
        width_factor = 0.72
    else:
        width_factor = 0.62
    single_line_min_size = 18.0 if ascii_ratio > 0.35 else 22.0
    target_lines = 1
    while font_size > single_line_min_size:
        capacity = max(1.0, box_width_pt / (font_size * width_factor))
        estimated_lines = math.ceil(text_units / capacity)
        if estimated_lines <= target_lines:
            return round(font_size, 2)
        font_size -= 1.0

    font_size = min(single_line_min_size, style.font_size_pt)
    two_line_min_size = 16.0 if ascii_ratio > 0.35 else 19.0
    while font_size > two_line_min_size:
        capacity = max(1.0, box_width_pt / (font_size * width_factor))
        estimated_lines = math.ceil(text_units / capacity)
        if estimated_lines <= 2:
            return round(font_size, 2)
        font_size -= 1.0
    return round(two_line_min_size, 2)


def add_header(slide, title_text: str, slide_w: int, slide_h: int, assets: TemplateAssets) -> dict[str, Any]:
    normalized_title = normalize_header_text(title_text)
    if not normalized_title:
        return {}
    style = choose_header_style(title_text, assets)
    header_width_ratio = choose_header_width_ratio(style, normalized_title)
    fitted_font_size = choose_fitted_header_font_size(style, slide_w, normalized_title, width_ratio=header_width_ratio)
    header = slide.shapes.add_textbox(
        int(slide_w * style.left_ratio),
        int(slide_h * style.top_ratio),
        int(slide_w * header_width_ratio),
        int(slide_h * style.height_ratio),
    )
    header.fill.background()
    header.line.fill.background()

    text_frame = header.text_frame
    text_frame.margin_left = style.margin_left
    text_frame.margin_right = style.margin_right
    text_frame.margin_top = style.margin_top
    text_frame.margin_bottom = style.margin_bottom
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = normalized_title
    run.font.bold = True
    run.font.size = Pt(fitted_font_size)
    run.font.color.rgb = rgb(style.color_hex)
    apply_font_family_to_run(run, style.font_name)

    return {
        "font_name": style.font_name,
        "font_size_pt": fitted_font_size,
        "color_hex": style.color_hex,
        "style": "wide" if style == assets.header_style_wide else "compact",
        "width_ratio": header_width_ratio,
    }


def style_cover_slide(slide, prs: Presentation, detected_title: str, palette: TemplatePalette, profile: TemplateProfile, assets: TemplateAssets) -> dict[str, Any]:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(palette.background_dark)
    background_applied = apply_template_background(slide, prs.slide_width, prs.slide_height, assets.cover_background_blob)
    text_hex = palette.text_dark if palette_background_is_light(palette) else palette.text_light
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and safe_text(shape):
            recolor_text_shape(shape, text_hex, font_name=palette.body_font)
            apply_text_autofit(shape, prs.slide_width, prs.slide_height)
    return {"mode": "cover_preview", "background_applied": background_applied, "title_kept": bool(detected_title)}


def style_section_slide(slide, prs: Presentation, slide_analysis: dict[str, Any], palette: TemplatePalette, assets: TemplateAssets) -> dict[str, Any]:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(palette.background_dark)
    background_applied = apply_template_background(slide, prs.slide_width, prs.slide_height, assets.section_background_blob)
    text_hex = palette.text_dark if palette_background_is_light(palette) else palette.text_light
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and safe_text(shape):
            recolor_text_shape(shape, text_hex, font_name=palette.body_font)
            apply_text_autofit(shape, prs.slide_width, prs.slide_height)
    return {"mode": "section_preview", "background_applied": background_applied}


def style_content_slide(
    slide,
    prs: Presentation,
    slide_analysis: dict[str, Any],
    palette: TemplatePalette,
    profile: TemplateProfile,
    assets: TemplateAssets,
) -> dict[str, Any]:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(palette.background_dark)
    slide_bg_luma = luminance_from_rgb_tuple(hex_to_rgb_tuple(palette.background_dark))
    safe_zone_bottom = int(prs.slide_height * profile.safe_zone_bottom_ratio)

    title_shape_index = slide_analysis["title_detection"]["shape_index"]
    title_text = slide_analysis["title_detection"]["title"]
    moved_shapes = []
    styled_tables = 0
    styled_charts = 0
    styled_groups = 0
    styled_vectors = 0
    original_shapes = list(slide.shapes)
    removed_title = False

    for shape_idx, shape in enumerate(original_shapes, start=1):
        if shape_idx == title_shape_index:
            text = safe_text(shape)
            if text:
                remove_shape(shape)
                removed_title = True
                continue
        if getattr(shape, "has_table", False):
            style_table_shape(shape, palette)
            styled_tables += 1
            continue
        if getattr(shape, "has_chart", False):
            style_chart_shape(shape, palette)
            styled_charts += 1
            continue
        if getattr(shape.shape_type, "name", str(shape.shape_type)) == "GROUP":
            group_meta = style_group_shape(shape, palette, prs.slide_width, prs.slide_height, slide_bg_luma)
            styled_groups += 1
            styled_tables += group_meta["tables"]
            styled_charts += group_meta["charts"]
            styled_vectors += group_meta["vectors"]
            continue

        text = safe_text(shape)
        if not text:
            vector_meta = style_vector_shape(
                shape,
                palette,
                prs.slide_width,
                prs.slide_height,
                slide_bg_luma,
                original_shapes[: shape_idx - 1],
            )
            if vector_meta:
                styled_vectors += 1
            continue
        if shape.top < safe_zone_bottom and shape.left < prs.slide_width * 0.62:
            delta = safe_zone_bottom - shape.top + int(prs.slide_height * 0.015)
            shift_shape_down(shape, delta, prs.slide_height)
            moved_shapes.append({"text": text[:60], "delta_emu": int(delta)})
        if is_nested_text_panel(shape, original_shapes[: shape_idx - 1]):
            clear_shape_fill_and_line(shape)
        light_fill = shape_local_bg_is_light(shape, slide_bg_luma, original_shapes[: shape_idx - 1])
        recolor_text_shape(
            shape,
            palette.text_dark if light_fill else palette.text_light,
            font_name=palette.body_font,
        )
        apply_text_autofit(shape, prs.slide_width, prs.slide_height)
        vector_meta = style_vector_shape(
            shape,
            palette,
            prs.slide_width,
            prs.slide_height,
            slide_bg_luma,
            original_shapes[: shape_idx - 1],
        )
        if vector_meta:
            styled_vectors += 1

    background_applied = apply_template_background(slide, prs.slide_width, prs.slide_height, assets.content_background_blob)
    header_meta = add_header(slide, title_text, prs.slide_width, prs.slide_height, assets)
    return {
        "mode": "content_preview",
        "background_applied": background_applied,
        "title_removed": removed_title,
        "header_added": bool(header_meta),
        "header_meta": header_meta,
        "moved_shapes": moved_shapes,
        "styled_tables": styled_tables,
        "styled_charts": styled_charts,
        "styled_groups": styled_groups,
        "styled_vectors": styled_vectors,
        "safe_zone_bottom_ratio": profile.safe_zone_bottom_ratio,
    }


def transform_deck(
    source_path: Path,
    output_pptx: Path,
    analysis: dict[str, Any],
    palette: TemplatePalette,
    profile: TemplateProfile,
    template_path: Path,
) -> dict[str, Any]:
    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_pptx)
    prs = Presentation(output_pptx)
    assets = extract_template_assets(template_path)
    transforms = []
    for slide_analysis in analysis["slides"]:
        slide = prs.slides[slide_analysis["slide_index"] - 1]
        role = slide_analysis["role"]
        if role == "cover":
            transform_meta = style_cover_slide(slide, prs, slide_analysis["title_detection"]["title"], palette, profile, assets)
        elif role == "section":
            transform_meta = style_section_slide(slide, prs, slide_analysis, palette, assets)
        else:
            transform_meta = style_content_slide(slide, prs, slide_analysis, palette, profile, assets)
        transforms.append({"slide_index": slide_analysis["slide_index"], "transform": transform_meta})
    prs.save(output_pptx)
    return {"output_pptx": str(output_pptx), "transforms": transforms, "template_profile": profile.__dict__}
