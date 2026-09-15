import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from prototype_tool.config import TemplatePalette
from prototype_tool.template_assets import HeaderStyle
from prototype_tool.transform import (
    apply_font_family_to_run,
    apply_text_autofit,
    choose_text_hex_for_fill,
    choose_fitted_header_font_size,
    is_nested_text_panel,
    style_vector_shape,
)


class TransformTests(unittest.TestCase):
    def test_apply_font_family_sets_latin_and_east_asian_nodes(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, 1_000_000, 300_000)
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "中文 Agent"

        apply_font_family_to_run(run, "腾讯体 W7")

        r_pr = run._r.get_or_add_rPr()
        self.assertEqual(r_pr.find(qn("a:latin")).get("typeface"), "腾讯体 W7")
        self.assertEqual(r_pr.find(qn("a:ea")).get("typeface"), "腾讯体 W7")
        self.assertEqual(r_pr.find(qn("a:cs")).get("typeface"), "腾讯体 W7")

    def test_choose_fitted_header_font_size_shrinks_long_mixed_title(self) -> None:
        style = HeaderStyle(
            left_ratio=0.07,
            top_ratio=0.04,
            width_ratio=0.61,
            height_ratio=0.10,
            margin_left=0,
            margin_right=0,
            margin_top=0,
            margin_bottom=0,
            font_name="腾讯体 W7",
            font_size_pt=32.0,
            color_hex="42D7FF",
        )
        short_size = choose_fitted_header_font_size(style, 9_144_000, "行业背景")
        long_size = choose_fitted_header_font_size(style, 9_144_000, "Successful Integration Case with Tanka")

        self.assertGreaterEqual(short_size, 30.0)
        self.assertLess(long_size, 32.0)
        self.assertGreaterEqual(long_size, 17.0)

    def test_apply_text_autofit_marks_small_text_boxes(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, int(prs.slide_width * 0.25), int(prs.slide_height * 0.08))
        textbox.text = "自动化程度提升和应用落地"

        apply_text_autofit(textbox, prs.slide_width, prs.slide_height)

        self.assertEqual(textbox.text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)

    def test_nested_text_panel_detection_prefers_outer_card(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        outer = slide.shapes.add_shape(1, 0, 0, 4_000_000, 2_000_000)
        outer.fill.solid()
        outer.fill.fore_color.rgb = RGBColor(0x11, 0x18, 0x27)
        inner = slide.shapes.add_textbox(200_000, 200_000, 3_600_000, 600_000)
        inner.text = "软件开发持续提效增质是企业不断追求的目标"

        self.assertTrue(is_nested_text_panel(inner, [outer]))

    def test_choose_text_hex_for_fill_uses_dark_text_on_light_surfaces(self) -> None:
        palette = TemplatePalette(text_light="F8FBFF", text_dark="1E293B")

        self.assertEqual(choose_text_hex_for_fill("EEF4FB", palette), palette.text_dark)
        self.assertEqual(choose_text_hex_for_fill("243B5A", palette), palette.text_light)

    def test_style_vector_shape_keeps_text_dark_on_light_card(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        outer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0, 0, 4_800_000, 2_200_000)
        outer.fill.solid()
        outer.fill.fore_color.rgb = RGBColor(0xEC, 0xF3, 0xFA)
        outer.line.color.rgb = RGBColor(0x4B, 0x70, 0x98)

        inner = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            250_000,
            250_000,
            3_600_000,
            700_000,
        )
        inner.fill.background()
        inner.line.fill.background()
        inner.text_frame.text = "配置：DATABASE + INSERT:100 + SEQUENTIAL"

        palette = TemplatePalette(
            background_dark="DDEFFA",
            text_light="F8FBFF",
            text_dark="1E293B",
            panel_mid="F5FAFE",
            panel_alt="EAF2FB",
            border_muted="5C7694",
        )

        style_vector_shape(
            inner,
            palette,
            prs.slide_width,
            prs.slide_height,
            0.9,
            [outer],
        )

        run = inner.text_frame.paragraphs[0].runs[0]
        self.assertEqual(str(run.font.color.rgb), palette.text_dark)


if __name__ == "__main__":
    unittest.main()
